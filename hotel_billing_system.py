import os
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl

EXCEL_FILE = "hotel_billing.xlsx"

# ----------------- 1. Excel Backend Operations -----------------
def init_excel():
    """Initializes the Excel workbook with headers if it does not already exist."""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bills"
        headers = [
            "Bill ID", "Customer Name", "Room Type", 
            "Days Stayed", "Daily Rate", "Food Charges", "Total Bill"
        ]
        ws.append(headers)
        wb.save(EXCEL_FILE)

def fetch_records():
    """Reads and returns all stored billing records from the Excel worksheet."""
    if not os.path.exists(EXCEL_FILE):
        init_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Bills"]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            records.append(list(row))
    wb.close()
    return records

def save_record(data):
    """Appends a new billing entry to the Excel file."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Bills"]
    ws.append(data)
    wb.save(EXCEL_FILE)

def update_record_in_excel(bill_id, updated_data):
    """Locates and updates an existing record matching the provided Bill ID."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Bills"]
    found = False
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(bill_id):
            for col_idx, val in enumerate(updated_data):
                row[col_idx].value = val
            found = True
            break
    wb.save(EXCEL_FILE)
    return found

def delete_record_from_excel(bill_id):
    """Deletes the row matching the specified Bill ID from the Excel sheet."""
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["Bills"]
    row_to_delete = None
    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if str(row[0].value) == str(bill_id):
            row_to_delete = idx
            break
    if row_to_delete:
        ws.delete_rows(row_to_delete, 1)
        wb.save(EXCEL_FILE)
        return True
    wb.close()
    return False

# ----------------- 2. Main Mobile Application Setup -----------------
class HotelBillingApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Hotel Billing System")

        init_excel()

        # Dynamically detect mobile screen resolution
        screen_w = self.winfo_screenwidth()
        screen_h = self.winfo_screenheight()
        self.geometry(f"{screen_w}x{screen_h}")
        self.resizable(True, True)

        # Responsive grid weights
        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=1)

        self.frames = {}
        for PageClass in (LoginPage, DashboardPage):
            frame = PageClass(parent=self, controller=self)
            self.frames[PageClass] = frame
            frame.grid(row=0, column=0, sticky="nsew")

        self.show_frame(LoginPage)

    def show_frame(self, page_class):
        """Switches the active view to the specified frame class."""
        frame = self.frames[page_class]
        if hasattr(frame, "on_show"):
            frame.on_show()
        frame.tkraise()

# ----------------- 3. Login Page -----------------
class LoginPage(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg="#2C3E50")
        self.controller = controller

        # Compact center card suitable for mobile displays
        card = tk.Frame(self, bg="#FFFFFF", bd=2, relief="groove", padx=20, pady=20)
        card.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(card, text="Hotel Billing", font=("Arial", 14, "bold"), bg="#FFFFFF", fg="#2C3E50").pack(pady=(0, 4))
        tk.Label(card, text="Admin Login", font=("Arial", 10), bg="#FFFFFF", fg="#7F8C8D").pack(pady=(0, 12))

        tk.Label(card, text="Username:", font=("Arial", 9, "bold"), bg="#FFFFFF").pack(anchor="w")
        self.entry_user = tk.Entry(card, font=("Arial", 10), width=20)
        self.entry_user.pack(pady=(2, 8))

        tk.Label(card, text="Password:", font=("Arial", 9, "bold"), bg="#FFFFFF").pack(anchor="w")
        self.entry_pass = tk.Entry(card, font=("Arial", 10), show="*", width=20)
        self.entry_pass.pack(pady=(2, 14))

        btn_login = tk.Button(card, text="Login", bg="#27AE60", fg="white", font=("Arial", 10, "bold"),
                              width=16, relief="flat", command=self.validate_login)
        btn_login.pack()

    def validate_login(self):
        """Validates administrator credentials."""
        username = self.entry_user.get().strip()
        password = self.entry_pass.get().strip()

        if username == "admin" and password == "admin123":
            self.entry_user.delete(0, tk.END)
            self.entry_pass.delete(0, tk.END)
            self.controller.show_frame(DashboardPage)
        else:
            messagebox.showerror("Error", "Invalid Login! (Use: admin / admin123)")

# ----------------- 4. Mobile Dashboard Page -----------------
class DashboardPage(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent, bg="#F4F6F7")
        self.controller = controller

        # Top navigation header
        header = tk.Frame(self, bg="#2980B9", height=40, padx=8)
        header.pack(fill="x", side="top")

        lbl_title = tk.Label(header, text="Hotel Billing Dashboard", font=("Arial", 11, "bold"), bg="#2980B9", fg="white")
        lbl_title.pack(side="left", pady=6)

        btn_logout = tk.Button(header, text="Logout", font=("Arial", 8, "bold"), bg="#E74C3C", 
                               fg="white", relief="flat", padx=6, command=self.logout)
        btn_logout.pack(side="right", pady=6)

        # Tab navigation for compact vertical mobile screens
        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=4, pady=4)

        self.tab_form = tk.Frame(self.notebook, bg="#FFFFFF", padx=8, pady=8)
        self.tab_view = tk.Frame(self.notebook, bg="#FFFFFF", padx=4, pady=4)

        self.notebook.add(self.tab_form, text="  1. Form Entry  ")
        self.notebook.add(self.tab_view, text="  2. Records & Search  ")

        self.setup_form_tab()
        self.setup_view_tab()

    def setup_form_tab(self):
        """Constructs input form fields and billing action buttons."""
        labels = [
            ("Bill ID:", "Bill ID"),
            ("Customer Name:", "Customer Name"),
            ("Room Type:", "Room Type"),
            ("Days Stayed:", "Days Stayed"),
            ("Daily Rate (Rs):", "Daily Rate"),
            ("Food Charges (Rs):", "Food Charges")
        ]

        self.entries = {}
        for row_idx, (lbl_text, key) in enumerate(labels):
            tk.Label(self.tab_form, text=lbl_text, font=("Arial", 9, "bold"), bg="#FFFFFF").grid(
                row=row_idx, column=0, sticky="w", pady=4, padx=2
            )

            if key == "Room Type":
                self.room_type_cb = ttk.Combobox(
                    self.tab_form, 
                    values=["Standard", "Deluxe", "Super Deluxe", "Suite"], 
                    state="readonly", 
                    width=15
                )
                self.room_type_cb.current(0)
                self.room_type_cb.grid(row=row_idx, column=1, pady=4, sticky="ew")
                self.room_type_cb.bind("<<ComboboxSelected>>", self.auto_fill_rate)
            else:
                entry = tk.Entry(self.tab_form, font=("Arial", 9), width=16)
                entry.grid(row=row_idx, column=1, pady=4, sticky="ew")
                self.entries[key] = entry

        self.entries["Daily Rate"].insert(0, "1000")
        self.entries["Food Charges"].insert(0, "0")

        # 2x2 Action Button Grid
        btn_frame = tk.Frame(self.tab_form, bg="#FFFFFF", pady=10)
        btn_frame.grid(row=len(labels), column=0, columnspan=2)

        tk.Button(btn_frame, text="Add Bill", bg="#27AE60", fg="white", font=("Arial", 8, "bold"),
                  width=10, pady=2, command=self.add_bill).grid(row=0, column=0, padx=4, pady=3)
        tk.Button(btn_frame, text="Update", bg="#F39C12", fg="white", font=("Arial", 8, "bold"),
                  width=10, pady=2, command=self.update_bill).grid(row=0, column=1, padx=4, pady=3)
        tk.Button(btn_frame, text="Delete", bg="#C0392B", fg="white", font=("Arial", 8, "bold"),
                  width=10, pady=2, command=self.delete_bill).grid(row=1, column=0, padx=4, pady=3)
        tk.Button(btn_frame, text="Clear Form", bg="#7F8C8D", fg="white", font=("Arial", 8, "bold"),
                  width=10, pady=2, command=self.clear_form).grid(row=1, column=1, padx=4, pady=3)

    def setup_view_tab(self):
        """Constructs search controls and scrollable records table."""
        # Search panel
        search_frame = tk.Frame(self.tab_view, bg="#FFFFFF", pady=4)
        search_frame.pack(fill="x")

        tk.Label(search_frame, text="Search:", font=("Arial", 8, "bold"), bg="#FFFFFF").pack(side="left")
        self.entry_search = tk.Entry(search_frame, font=("Arial", 8), width=12)
        self.entry_search.pack(side="left", padx=3)

        tk.Button(search_frame, text="Go", bg="#3498DB", fg="white", font=("Arial", 8, "bold"),
                  command=self.search_record).pack(side="left", padx=2)
        tk.Button(search_frame, text="Reset", bg="#95A5A6", fg="white", font=("Arial", 8),
                  command=self.load_table_data).pack(side="left", padx=2)

        # Treeview table container with horizontal and vertical scrollbars
        table_container = tk.Frame(self.tab_view)
        table_container.pack(fill="both", expand=True, pady=4)

        columns = ("Bill ID", "Customer Name", "Room Type", "Days", "Rate", "Food", "Total")
        self.tree = ttk.Treeview(table_container, columns=columns, show="headings", height=8)

        col_widths = {
            "Bill ID": 60, "Customer Name": 110, "Room Type": 85, 
            "Days": 45, "Rate": 55, "Food": 55, "Total": 75
        }
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=col_widths.get(col, 60), anchor="center")
        self.tree.column("Customer Name", anchor="w")

        scroll_y = ttk.Scrollbar(table_container, orient="vertical", command=self.tree.yview)
        scroll_x = ttk.Scrollbar(table_container, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)

        scroll_y.pack(side="right", fill="y")
        scroll_x.pack(side="bottom", fill="x")
        self.tree.pack(side="left", fill="both", expand=True)

        self.tree.bind("<<TreeviewSelect>>", self.on_row_select)

    def on_show(self):
        """Refreshes records whenever the dashboard becomes visible."""
        self.load_table_data()

    def auto_fill_rate(self, event=None):
        """Updates default rate according to chosen room category."""
        rates = {"Standard": "1000", "Deluxe": "2000", "Super Deluxe": "3500", "Suite": "5000"}
        selected = self.room_type_cb.get()
        self.entries["Daily Rate"].delete(0, tk.END)
        self.entries["Daily Rate"].insert(0, rates.get(selected, "1000"))

    def compute_total(self, days, rate, food):
        """Calculates the total billing amount."""
        return (days * rate) + food

    def load_table_data(self, records=None):
        """Reloads treeview contents from records list or directly from Excel."""
        for item in self.tree.get_children():
            self.tree.delete(item)
        if records is None:
            records = fetch_records()
        for rec in records:
            self.tree.insert("", tk.END, values=rec)

    def add_bill(self):
        """Validates inputs and appends new bill entry to Excel."""
        bill_id = self.entries["Bill ID"].get().strip()
        cust_name = self.entries["Customer Name"].get().strip()
        room_type = self.room_type_cb.get()

        if not bill_id or not cust_name:
            messagebox.showwarning("Warning", "Bill ID and Customer Name are required!")
            return

        try:
            days = int(self.entries["Days Stayed"].get().strip())
            rate = float(self.entries["Daily Rate"].get().strip())
            food = float(self.entries["Food Charges"].get().strip())
            if days <= 0 or rate < 0 or food < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Error", "Days, Rate, and Food Charges must be valid numbers.")
            return

        # Check for unique Bill ID
        for row in fetch_records():
            if str(row[0]) == bill_id:
                messagebox.showerror("Duplicate", f"Bill ID '{bill_id}' already exists!")
                return

        total = self.compute_total(days, rate, food)
        data = [bill_id, cust_name, room_type, days, rate, food, total]

        save_record(data)
        messagebox.showinfo("Success", f"Bill saved successfully!\nTotal: Rs {total:.2f}")
        self.load_table_data()
        self.clear_form()
        self.notebook.select(self.tab_view)

    def on_row_select(self, event):
        """Loads selected row values into form entries and switches to Form tab."""
        selected = self.tree.focus()
        if not selected:
            return
        values = self.tree.item(selected, "values")
        if values:
            self.clear_form()
            self.entries["Bill ID"].insert(0, values[0])
            self.entries["Customer Name"].insert(0, values[1])
            self.room_type_cb.set(values[2])
            self.entries["Days Stayed"].insert(0, values[3])
            self.entries["Daily Rate"].insert(0, values[4])
            self.entries["Food Charges"].insert(0, values[5])
            self.notebook.select(self.tab_form)

    def update_bill(self):
        """Updates record in Excel matching current Bill ID."""
        bill_id = self.entries["Bill ID"].get().strip()
        if not bill_id:
            messagebox.showwarning("Warning", "Select a record or enter Bill ID to update.")
            return

        try:
            days = int(self.entries["Days Stayed"].get().strip())
            rate = float(self.entries["Daily Rate"].get().strip())
            food = float(self.entries["Food Charges"].get().strip())
            cust_name = self.entries["Customer Name"].get().strip()
            room_type = self.room_type_cb.get()
            total = self.compute_total(days, rate, food)
        except ValueError:
            messagebox.showerror("Error", "Please provide valid numeric values.")
            return

        updated_data = [bill_id, cust_name, room_type, days, rate, food, total]
        if update_record_in_excel(bill_id, updated_data):
            messagebox.showinfo("Success", f"Bill ID {bill_id} updated successfully.")
            self.load_table_data()
            self.clear_form()
            self.notebook.select(self.tab_view)
        else:
            messagebox.showerror("Error", f"Bill ID {bill_id} not found.")

    def delete_bill(self):
        """Deletes record from Excel after user confirmation."""
        bill_id = self.entries["Bill ID"].get().strip()
        if not bill_id:
            messagebox.showwarning("Warning", "Select a record or enter Bill ID to delete.")
            return

        confirm = messagebox.askyesno("Confirm", f"Are you sure you want to delete Bill ID: {bill_id}?")
        if confirm:
            if delete_record_from_excel(bill_id):
                messagebox.showinfo("Deleted", f"Bill ID {bill_id} has been removed.")
                self.load_table_data()
                self.clear_form()
            else:
                messagebox.showerror("Error", f"Bill ID {bill_id} not found.")

    def clear_form(self):
        """Resets all input fields back to their default values."""
        for entry in self.entries.values():
            entry.delete(0, tk.END)
        self.room_type_cb.current(0)
        self.auto_fill_rate()
        self.entries["Food Charges"].insert(0, "0")

    def search_record(self):
        """Filters displayed table records matching the search query."""
        query = self.entry_search.get().strip().lower()
        if not query:
            self.load_table_data()
            return
        records = fetch_records()
        filtered = [r for r in records if query in str(r[0]).lower() or query in str(r[1]).lower()]
        self.load_table_data(filtered)

    def logout(self):
        """Logs out of the dashboard session and returns to login page."""
        confirm = messagebox.askyesno("Logout", "Do you want to log out?")
        if confirm:
            self.clear_form()
            self.controller.show_frame(LoginPage)

# ----------------- 5. Application Execution -----------------
if __name__ == "__main__":
    app = HotelBillingApp()
    app.mainloop()
